"""Reopen actual outputs independently; render PDF samples with PDFium."""
from pathlib import Path
import json
import openpyxl
import pypdfium2 as pdfium
from pypdf import PdfReader
root=Path('outputs/radical')
report=[]
for name,expected in [('daily-source.private',196),('monthly-source.private',6360),('monthly-test-window.private',424)]:
    wb=openpyxl.load_workbook(root/(name+'.xlsx'),data_only=False)
    ws=wb.active
    rows=list(ws.values)
    assert ws.sheet_view.rightToLeft
    assert ws.freeze_panes and ws.auto_filter.ref
    assert len(ws._images)==2
    assert ws.max_row==expected+5
    assert sum(1 for row in ws if row[0].row>5 and row[0].value is not None)==expected
    assert not any(c.data_type=='f' for row in ws for c in row)
    reader=PdfReader(root/(name+'.pdf'))
    doc=pdfium.PdfDocument(root/(name+'.pdf'))
    pages=len(doc)
    sample=sorted(set([0,pages//2,pages-1]))
    for page in sample:
        doc[page].render(scale=1.2).to_pil().save(root/(name+f'-page-{page+1}.png'))
    fonts=reader.pages[0]['/Resources']['/Font']
    assert any('/FontFile2' in str(v.get_object()) or '/Type0' in str(v.get_object()) for v in fonts.values())
    report.append({'file':name,'excel_rows':ws.max_row,'data_rows':expected,'columns':ws.max_column,'rtl':True,'freeze_panes':ws.freeze_panes,'filter':ws.auto_filter.ref,'images':len(ws._images),'pdf_pages':pages,'rendered_pages':[p+1 for p in sample]})
(root/'export-inspection.json').write_text(json.dumps(report,indent=2),encoding='utf8')
print(json.dumps(report,indent=2))
